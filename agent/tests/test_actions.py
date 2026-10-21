"""Unit and integration tests for all 7 core Flash actions."""

from pathlib import Path
import openpyxl
from fastapi.testclient import TestClient


def test_action_create_text_file(client: TestClient) -> None:
    """Action 3: create_text_file."""
    resp = client.post(
        "/api/v1/actions/execute",
        json={
            "action": "create_text_file",
            "parameters": {
                "filename": "documents/hello_flash.txt",
                "content": "Hello world from Flash the Sloth!\nFast as lightning...",
                "overwrite": True,
            },
        },
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert data["action"] == "create_text_file"
    output_path = Path(data["output_path"])
    assert output_path.exists()
    assert output_path.read_text(encoding="utf-8").startswith("Hello world from Flash")
    assert data["data"]["lines"] == 2


def test_action_create_grocery_list(client: TestClient) -> None:
    """Action 4: create_grocery_list."""
    items = [
        {"category": "Fresh Produce", "items": ["Avocados", "Eucalyptus leaves", "Bananas"]},
        {"category": "Bakery & Pantry", "items": ["Sourdough bread", "Olive oil"]},
    ]
    resp = client.post(
        "/api/v1/actions/execute",
        json={
            "action": "create_grocery_list",
            "parameters": {
                "title": "Weekend Sloth Feast",
                "items": items,
                "filename": "lists/sloth_feast.md",
            },
        },
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert data["action"] == "create_grocery_list"
    assert data["data"]["item_count"] == 5

    output_path = Path(data["output_path"])
    assert output_path.exists()
    content = output_path.read_text(encoding="utf-8")
    assert "- [ ] Avocados" in content
    assert "### Fresh Produce" in content
    assert "Weekend Sloth Feast" in content


def test_action_create_spreadsheet(client: TestClient) -> None:
    """Action 5: create_spreadsheet - real openpyxl .xlsx generation."""
    headers = ["Task", "Category", "Hours", "Completed"]
    rows = [
        ["Inspect DMV licenses", "Administration", 4.5, "Yes"],
        ["Coffee break", "Wellness", 1.5, "Yes"],
        ["Review automation logs", "Engineering", 3.0, "No"],
    ]

    resp = client.post(
        "/api/v1/actions/execute",
        json={
            "action": "create_spreadsheet",
            "parameters": {
                "headers": headers,
                "rows": rows,
                "filename": "spreadsheets/sloth_timesheet.xlsx",
                "sheet_name": "Timesheet",
                "include_total_row": True,
            },
        },
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert data["action"] == "create_spreadsheet"

    output_path = Path(data["output_path"])
    assert output_path.exists()
    assert output_path.suffix.lower() == ".xlsx"

    # Verify the created Excel file with openpyxl
    wb = openpyxl.load_workbook(output_path, data_only=False)
    assert "Timesheet" in wb.sheetnames
    ws = wb["Timesheet"]

    # Verify headers
    assert ws.cell(row=1, column=1).value == "Task"
    assert ws.cell(row=1, column=3).value == "Hours"

    # Verify data row
    assert ws.cell(row=2, column=1).value == "Inspect DMV licenses"
    assert ws.cell(row=2, column=3).value == 4.5

    # Verify formula row
    assert ws.cell(row=5, column=1).value == "Total"
    assert ws.cell(row=5, column=3).value == "=SUM(C2:C4)"


def test_action_open_file(client: TestClient) -> None:
    """Action 6: open_file."""
    # First create a file to open
    create_resp = client.post(
        "/api/v1/actions/execute",
        json={
            "action": "create_text_file",
            "parameters": {"filename": "documents/open_me.txt", "content": "Flash authorized file."},
        },
    )
    assert create_resp.status_code == 200
    created_path = create_resp.json()["output_path"]

    # Now open it via open_file
    open_resp = client.post(
        "/api/v1/actions/execute",
        json={"action": "open_file", "parameters": {"filepath": "documents/open_me.txt"}},
    )
    assert open_resp.status_code == 200
    data = open_resp.json()
    assert data["success"] is True
    assert data["action"] == "open_file"
    assert data["output_path"] == created_path


def test_action_open_file_not_found(client: TestClient) -> None:
    """Verify open_file fails gracefully if file does not exist."""
    resp = client.post(
        "/api/v1/actions/execute",
        json={"action": "open_file", "parameters": {"filepath": "non_existent.txt"}},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is False
    assert "PATH_ACCESS_DENIED" in data["error"] or "does not exist" in data["message"].lower()


def test_action_open_url_valid(client: TestClient, monkeypatch) -> None:
    """Action 1: open_url - safe URL dispatch."""
    import webbrowser

    # Mock actual browser spawn to avoid opening real popups during automated test runs
    opened_urls = []
    monkeypatch.setattr(webbrowser, "open", lambda url: opened_urls.append(url) or True)

    resp = client.post(
        "/api/v1/actions/execute",
        json={"action": "open_url", "parameters": {"url": "https://fastapi.tiangolo.com"}},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert data["action"] == "open_url"
    assert "https://fastapi.tiangolo.com" in opened_urls


def test_action_open_url_invalid(client: TestClient) -> None:
    """Action 1: open_url - blocks invalid URLs."""
    resp = client.post(
        "/api/v1/actions/execute",
        json={"action": "open_url", "parameters": {"url": "javascript:alert(1)"}},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is False
    assert data["error"] == "UNSAFE_URL"


def test_action_open_url_with_query_parameter(client: TestClient, monkeypatch) -> None:
    """Action 1: open_url - supports 'query' parameter fallback."""
    import webbrowser

    opened_urls = []
    monkeypatch.setattr(webbrowser, "open", lambda url: opened_urls.append(url) or True)

    resp = client.post(
        "/api/v1/actions/execute",
        json={"action": "open_url", "parameters": {"query": "something"}},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert "https://www.google.com/search?q=something" in opened_urls



def test_action_type_text(client: TestClient) -> None:
    """Action 2: type_text - verify validation and parameter handling."""
    resp = client.post(
        "/api/v1/actions/execute",
        json={"action": "type_text", "parameters": {"text": "Flash is typing...", "interval_ms": 0, "press_return": False}},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert data["action"] == "type_text"
    assert data["data"]["characters_typed"] == 18


def test_action_type_text_missing_param(client: TestClient) -> None:
    """Action 2: type_text - error when missing required parameter."""
    resp = client.post(
        "/api/v1/actions/execute",
        json={"action": "type_text", "parameters": {}},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is False
    assert data["error"] == "MISSING_PARAMETER"


def test_action_get_agent_status(client: TestClient) -> None:
    """Action 7: get_agent_status."""
    resp = client.post(
        "/api/v1/actions/execute",
        json={"action": "get_agent_status", "parameters": {}},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert data["action"] == "get_agent_status"
    assert data["data"]["character"] == "Flash"
    assert data["data"]["status"] == "ready"
    assert data["data"]["action_count"] >= 7
